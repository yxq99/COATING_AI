"""根据项目设置生成适合打印或填写的初始实验Excel表。"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from coating_ai.config import DATA_DIR, DEFAULT_SETTINGS_PATH, load_project_config
from coating_ai.data import frame, initial_points


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="生成初始实验Excel表")
    parser.add_argument("--settings", default=str(DEFAULT_SETTINGS_PATH))
    parser.add_argument("--output", help="输出xlsx路径")
    parser.add_argument("--n", type=int, help="覆盖设置中的初始实验数量")
    parser.add_argument("--seed", type=int, help="覆盖设置中的随机种子")
    return parser


def main() -> None:
    args = build_parser().parse_args()
    config = load_project_config(args.settings)
    count = args.n or config.initial_design_size
    if count < 1:
        raise ValueError("n必须大于0。")
    output = Path(args.output) if args.output else DATA_DIR / "initial_design.xlsx"
    if output.exists():
        raise FileExistsError(f"为防止覆盖已有记录，输出文件已经存在：{output}")
    output.parent.mkdir(parents=True, exist_ok=True)

    table = frame(config, initial_points(config, count, args.seed))
    table.to_excel(output, index=False, engine="openpyxl")
    workbook = load_workbook(output)
    sheet = workbook.active
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column) + 2
        sheet.column_dimensions[column[0].column_letter].width = min(max(width, 12), 24)
    workbook.save(output)
    print(f"已生成{count}组初始实验Excel：{output.resolve()}")


if __name__ == "__main__":
    main()
